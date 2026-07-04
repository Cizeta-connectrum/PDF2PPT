import os
import sys
import tempfile

import fitz
import pytest
from pptx import Presentation

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

from converter import convert_pdf_to_pptx  # noqa: E402


@pytest.fixture
def sample_pdf(tmp_path):
    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    for i in range(3):
        page = doc.new_page(width=612, height=792)  # US Letter
        page.insert_text((72, 72), f"Page {i + 1} heading", fontsize=24)
        page.insert_text((72, 120), "Some body text to verify extraction.", fontsize=12)
        page.draw_rect(fitz.Rect(72, 200, 300, 260), color=(0, 0, 1), fill=(0.8, 0.9, 1))
    doc.save(str(pdf_path))
    doc.close()
    return str(pdf_path)


def test_image_mode(sample_pdf, tmp_path):
    out_path = tmp_path / "out_image.pptx"
    result = convert_pdf_to_pptx(sample_pdf, str(out_path), mode="image")
    assert result.page_count == 3
    assert out_path.exists()

    prs = Presentation(str(out_path))
    assert len(prs.slides.__iter__.__self__._sldIdLst) == 3
    for slide in prs.slides:
        pictures = [s for s in slide.shapes if s.shape_type == 13]
        assert len(pictures) == 1


def test_editable_mode(sample_pdf, tmp_path):
    out_path = tmp_path / "out_editable.pptx"
    result = convert_pdf_to_pptx(sample_pdf, str(out_path), mode="editable")
    assert result.page_count == 3
    assert out_path.exists()

    prs = Presentation(str(out_path))
    slide_count = sum(1 for _ in prs.slides)
    assert slide_count == 3

    first_slide = next(iter(prs.slides))
    texts = []
    for shape in first_slide.shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
    joined = " ".join(texts)
    assert "Page 1 heading" in joined
    assert "body text" in joined


def test_invalid_mode(sample_pdf, tmp_path):
    out_path = tmp_path / "out.pptx"
    with pytest.raises(ValueError):
        convert_pdf_to_pptx(sample_pdf, str(out_path), mode="bogus")


@pytest.fixture
def flattened_image_pdf(tmp_path):
    """A PDF page that is purely a raster image with no text objects at
    all, mimicking a screenshot/scan-style export."""
    import ocr as ocr_module

    if not ocr_module.is_available():
        pytest.skip("Tesseract OCR is not installed in this environment")

    src = fitz.open()
    page = src.new_page(width=640, height=360)
    page.draw_rect(fitz.Rect(0, 0, 640, 140), color=None, fill=(0.08, 0.24, 0.47))
    page.insert_text((40, 80), "Overview Slide", fontsize=28, color=(1, 1, 1))
    page.insert_text((40, 200), "Body text baked into a flat image.", fontsize=16, color=(0, 0, 0))
    pix = page.get_pixmap(matrix=fitz.Matrix(3, 3))
    img_bytes = pix.tobytes("png")
    src.close()

    out = fitz.open()
    flat_page = out.new_page(width=640, height=360)
    flat_page.insert_image(flat_page.rect, stream=img_bytes)
    pdf_path = tmp_path / "flat.pdf"
    out.save(str(pdf_path))
    out.close()
    return str(pdf_path)


def test_editable_mode_falls_back_to_ocr(flattened_image_pdf, tmp_path):
    out_path = tmp_path / "out_ocr.pptx"
    result = convert_pdf_to_pptx(flattened_image_pdf, str(out_path), mode="editable")
    assert result.page_count == 1
    assert out_path.exists()

    prs = Presentation(str(out_path))
    slide = next(iter(prs.slides))
    pictures = [s for s in slide.shapes if s.shape_type == 13]
    assert len(pictures) == 1

    texts = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Overview" in texts or "Slide" in texts
    assert "Body" in texts or "flat image" in texts
